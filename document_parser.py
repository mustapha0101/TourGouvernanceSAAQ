"""
Document Parser pour le Guichet d'Évaluation des Initiatives IA d'Investissement Québec.
Prend en charge : DOCX, PPTX, XLSX, PDF.
Extrait le texte brut, les métadonnées et la structure pour alimenter les agents d'analyse.
"""

import os
from typing import Dict, Any
import docx
from pptx import Presentation
import openpyxl
from pypdf import PdfReader

def parse_docx(file_path: str) -> Dict[str, Any]:
    doc = docx.Document(file_path)
    full_text = []
    
    # Paragraphes
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            full_text.append(txt)
            
    # Tableaux
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
                
    content = "\n".join(full_text)
    return {
        "file_type": "docx",
        "file_name": os.path.basename(file_path),
        "text": content,
        "length": len(content),
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables)
    }

def parse_pptx(file_path: str) -> Dict[str, Any]:
    prs = Presentation(file_path)
    full_text = []
    slide_count = len(prs.slides)
    
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    txt = paragraph.text.strip()
                    if txt:
                        slide_texts.append(txt)
        if slide_texts:
            full_text.append(f"--- Diapositive {i} ---")
            full_text.extend(slide_texts)
            
    content = "\n".join(full_text)
    return {
        "file_type": "pptx",
        "file_name": os.path.basename(file_path),
        "text": content,
        "length": len(content),
        "slides_count": slide_count
    }

def parse_xlsx(file_path: str) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    full_text = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        full_text.append(f"--- Feuille : {sheet_name} ---")
        for row in ws.iter_rows(values_only=True):
            row_str = [str(c) for c in row if c is not None and str(c).strip()]
            if row_str:
                full_text.append(" | ".join(row_str))
                
    content = "\n".join(full_text)
    return {
        "file_type": "xlsx",
        "file_name": os.path.basename(file_path),
        "text": content,
        "length": len(content),
        "sheets_count": len(wb.sheetnames)
    }

def parse_pdf(file_path: str) -> Dict[str, Any]:
    reader = PdfReader(file_path)
    full_text = []
    for i, page in enumerate(reader.pages, 1):
        txt = page.extract_text()
        if txt and txt.strip():
            full_text.append(f"--- Page {i} ---")
            full_text.append(txt.strip())
            
    content = "\n".join(full_text)
    return {
        "file_type": "pdf",
        "file_name": os.path.basename(file_path),
        "text": content,
        "length": len(content),
        "pages_count": len(reader.pages)
    }

def extract_document_content(file_path: str) -> Dict[str, Any]:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return parse_docx(file_path)
    elif ext == ".pptx":
        return parse_pptx(file_path)
    elif ext == ".xlsx" or ext == ".xls":
        return parse_xlsx(file_path)
    elif ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        return {
            "file_type": ext.lstrip("."),
            "file_name": os.path.basename(file_path),
            "text": content,
            "length": len(content)
        }
    else:
        raise ValueError(f"Format de document non supporté : {ext}. Formats acceptés : .docx, .pptx, .xlsx, .pdf, .txt, .md")
