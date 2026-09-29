import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE_DIR = "/Users/mustaphaberrabaa/Documents/IQ/InitiativesIA/12_Metropolis_Gouvernance_IA_3D"
DOCS_DIR = os.path.join(BASE_DIR, "documents_phases")

PHASES = [
    {
        "step": 0,
        "folder": "Etape_0_Depot",
        "title": "Dépôt de la demande",
        "files": [
            {
                "filename": "Formulaire_Depot_Demande_IA_Octopus.xlsx",
                "title": "Formulaire de Prise en Charge Octopus IA",
                "desc": "Gabarit de saisie initiale des besoins d'affaires et orientation vers l'un des 3 parcours officiels.",
                "type": "excel"
            },
            {
                "filename": "Gabarit_Expression_Besoin_Affaires.docx",
                "title": "Gabarit d'Expression de Besoin d'Affaires",
                "desc": "Fiche de cadrage sommaire du problème opérationnel et des gains anticipés.",
                "type": "word"
            }
        ]
    },
    {
        "step": 1,
        "folder": "Etape_1_Qualification",
        "title": "Qualification (Bureau de l'IA)",
        "files": [
            {
                "filename": "Grille_Evaluation_Valeur_Faisabilite_Bureau_IA.xlsx",
                "title": "Grille d'Évaluation Valeur & Faisabilité (Bureau IA)",
                "desc": "Outil d'analyse multicritère de la valeur stratégique, maturité des données et faisabilité technique.",
                "type": "excel"
            },
            {
                "filename": "Arbre_Decision_Outils_Homologues_vs_Nouveaux.xlsx",
                "title": "Matrice de Décision Outils Homologués vs Développement",
                "desc": "Filtre anti-redondance vérifiant la couverture par Copilot M365, Read AI ou solution existante.",
                "type": "excel"
            },
            {
                "filename": "Fiche_Cadrage_Preliminaire_Cas_Usage.docx",
                "title": "Fiche de Cadrage Préliminaire du Cas d'Usage",
                "desc": "Synthèse de recevabilité et recommandation d'orientation du Bureau de l'IA.",
                "type": "word"
            }
        ]
    },
    {
        "step": 2,
        "folder": "Etape_2_Risques_Priorisation",
        "title": "Évaluation des risques et priorisation",
        "files": [
            {
                "filename": "Grille_Appetit_Risque_11_Risques_DGIR_2026.xlsx",
                "title": "Grille d'Appétit au Risque DGIR 2026 (11 Risques)",
                "desc": "Évaluation matricielle des 11 risques de gouvernance IA et tolérances institutionnelles.",
                "type": "excel"
            },
            {
                "filename": "Questionnaire_Evaluation_Loi_25_EFVP.xlsx",
                "title": "Questionnaire Préliminaire ÉFVP (Loi 25)",
                "desc": "Checklist de conformité sur la protection des renseignements personnels et sensibilité des données.",
                "type": "excel"
            },
            {
                "filename": "Matrice_Priorisation_Multicritere_AHP_TOPSIS.xlsx",
                "title": "Matrice de Priorisation Portefeuille AHP-TOPSIS",
                "desc": "Calcul vectoriel éliminant les faux positifs pour le classement du portefeuille.",
                "type": "excel"
            }
        ]
    },
    {
        "step": 3,
        "folder": "Etape_3_Decision_GoNoGo",
        "title": "Décision Go / No Go (Comités)",
        "files": [
            {
                "filename": "Memo_Decisionnel_Comite_Gouvernance_IA.docx",
                "title": "Mémo Décisionnel pour le Comité IA",
                "desc": "Document officiel d'arbitrage soumis au Comité de Gouvernance IA ou au Comité de Direction.",
                "type": "word"
            },
            {
                "filename": "Registre_Arbitrages_Decisions_Comite.xlsx",
                "title": "Registre des Décisions & Seuils d'Arbitrage",
                "desc": "Tableau de bord de suivi des conditions d'homologation et mandats accordés.",
                "type": "excel"
            }
        ]
    },
    {
        "step": 4,
        "folder": "Etape_4_Pilote_POC",
        "title": "Développement encadré (Pilote)",
        "files": [
            {
                "filename": "Grille_Evaluation_Bilan_Pilote_POC.xlsx",
                "title": "Grille de Bilan Pilote POC & Métriques Valeur",
                "desc": "Mesure des gains réels en environnement Azure isolé et vérification du seuil 'Fail Fast'.",
                "type": "excel"
            },
            {
                "filename": "Cahier_Charges_Experimentation_Escouade_IA.docx",
                "title": "Cahier des Charges d'Expérimentation Escouade IA",
                "desc": "Périmètre du pilote, jeu de données de test et protocoles de supervision humaine.",
                "type": "word"
            }
        ]
    },
    {
        "step": 5,
        "folder": "Etape_5_Deploiement_Controle",
        "title": "Déploiement contrôlé",
        "files": [
            {
                "filename": "Gabarit_ARP_Homologation_Securite_CSI.xlsx",
                "title": "Gabarit ARP & Homologation Sécurité TI (CSI)",
                "desc": "Vérification des passerelles API, Entra ID et validation finale de cybersécurité.",
                "type": "excel"
            },
            {
                "filename": "Plan_Formation_Accompagnement_Changement.docx",
                "title": "Plan de Formation & Adoption Utilisateurs",
                "desc": "Stratégie de conduite du changement et guides d'usage responsable pour les équipes métiers.",
                "type": "word"
            }
        ]
    },
    {
        "step": 6,
        "folder": "Etape_6_Exploitation_Surveillance",
        "title": "Exploitation et surveillance",
        "files": [
            {
                "filename": "Tableau_Bord_Suivi_KRIs_Derive_Modeles.xlsx",
                "title": "Tableau de Bord 24/7 de Suivi des KRIs & Dérive",
                "desc": "Surveillance opérationnelle TI, alertes d'incidents Octopus et qualité continue des extrants.",
                "type": "excel"
            },
            {
                "filename": "Registre_Officiel_Initiatives_IA_Actives.xlsx",
                "title": "Registre Officiel des Initiatives IA Actives d'IQ",
                "desc": "Inventaire institutionnel consolidé avec propriétaires d'affaires et revues annuelles.",
                "type": "excel"
            }
        ]
    }
]

def create_excel_template(file_path, title, desc, step_title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gabarit Officiel IQ"
    
    # Couleurs institutionnelles IQ
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    blue_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    light_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=11, color="1E293B")
    meta_font = Font(name="Calibri", size=10, italic=True, color="475569")
    
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Titre bandeau
    ws.merge_cells("A1:G1")
    cell_a1 = ws["A1"]
    cell_a1.value = "INVESTISSEMENT QUÉBEC — DIRECTION DES TECHNOLOGIES & BUREAU DE L'IA"
    cell_a1.fill = navy_fill
    cell_a1.font = title_font
    cell_a1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Sous-titre
    ws.merge_cells("A2:G2")
    cell_a2 = ws["A2"]
    cell_a2.value = f"Cadre de Gouvernance IA • {step_title} • {title}"
    cell_a2.fill = blue_fill
    cell_a2.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    cell_a2.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 26

    # Description & Instructions
    ws["A4"] = "Description du Gabarit :"
    ws["A4"].font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    ws["B4"] = desc
    ws["B4"].font = regular_font

    ws["A5"] = "Statut :"
    ws["A5"].font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    ws["B5"] = "Gabarit de travail modifiable — Remplacez par vos indicateurs et critères d'évaluation."
    ws["B5"].font = meta_font

    # Tableau exemple
    headers = ["ID Critère", "Dimension d'Évaluation", "Pondération", "Échelle (1 à 5)", "Score Initié", "Commentaires & Mitigations", "Responsable"]
    ws.row_dimensions[7].height = 26
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin

    sample_rows = [
        ("CRT-01", "Alignement stratégique et valeur métier attendue", "25%", "4", "4.0", "Gain de productivité significatif validé par le gestionnaire.", "Porteur Métier"),
        ("CRT-02", "Faisabilité technique & maturité de l'architecture", "20%", "4", "3.5", "Intégration via passerelles sécurisées Azure IQ.", "Bureau de l'IA"),
        ("CRT-03", "Conformité Loi 25 & Protection des Renseignements Personnels", "30%", "5", "4.5", "Aucune donnée sensible client non chiffrée.", "DPRP / CSI"),
        ("CRT-04", "Gestion du changement & autonomie des équipes", "15%", "3", "3.0", "Ateliers de formation Escouade IA planifiés.", "Bureau de l'IA"),
        ("CRT-05", "Résilience opérationnelle & continuité de service", "10%", "4", "4.0", "RTO conforme aux exigences opérationnelles d'IQ.", "Équipe TI")
    ]

    for r_idx, row in enumerate(sample_rows, start=8):
        ws.row_dimensions[r_idx].height = 22
        for c_idx, val in enumerate(row, start=1):
            c = ws.cell(row=r_idx, column=c_idx, value=val)
            c.font = regular_font
            c.border = border_thin
            if c_idx in (1, 3, 4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
            if r_idx % 2 == 0:
                c.fill = light_fill

    # Largeurs de colonnes
    col_widths = [14, 45, 15, 16, 15, 50, 22]
    for idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(idx)
        ws.column_dimensions[col_letter].width = width

    wb.save(file_path)

def create_word_template(file_path, title, desc, step_title):
    doc = Document()
    
    # Titre Principal
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_after = Pt(4)
    run_org = p_org.add_run("INVESTISSEMENT QUÉBEC — BUREAU DE L'IA & TI")
    run_org.font.name = "Calibri"
    run_org.font.size = Pt(10)
    run_org.font.bold = True
    run_org.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(18)
    run_meta = p_meta.add_run(f"Cadre de Gouvernance IA d'Investissement Québec • {step_title}")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(11)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Encadré Description
    table_desc = doc.add_table(rows=1, cols=1)
    table_desc.style = 'Light Shading Accent 1'
    cell = table_desc.cell(0, 0)
    cell.paragraphs[0].text = f"Objet du Document :\n{desc}\n\nCe gabarit officiel est prêt pour la saisie des informations du cas d'usage."

    doc.add_paragraph()

    # Section 1
    h1 = doc.add_heading("1. Synthèse Exécutive de la Demande", level=1)
    doc.add_paragraph(
        "Ce document a pour but de consigner formellement les paramètres du cas d'usage conformément au "
        "Cadre Unifié de Gouvernance IA d'Investissement Québec (arrimage Modèle de Confiance Numérique du Québec et NIST AI RMF)."
    )

    # Section 2
    doc.add_heading("2. Paramètres Opérationnels & Porteur Mandaté", level=1)
    doc.add_paragraph("• Porteur Mandaté (Ligne d'Affaires) : [À préciser]")
    doc.add_paragraph("• Direction Engagée : [À préciser]")
    doc.add_paragraph("• Gain de Productivité Estimé : [À préciser en heures / an]")
    doc.add_paragraph("• Niveau d'Autonomie & Supervision Humaine : [Humain dans la boucle]")

    # Section 3
    doc.add_heading("3. Décision & Recommandation du Bureau de l'IA", level=1)
    doc.add_paragraph(
        "Après qualification selon l'arbre de décision officiel, la présente initiative est orientée vers le jalon suivant avec les conditions de suivi requises."
    )

    doc.save(file_path)

def main():
    os.makedirs(DOCS_DIR, exist_ok=True)
    manifest_data = []

    for phase in PHASES:
        phase_dir = os.path.join(DOCS_DIR, phase["folder"])
        os.makedirs(phase_dir, exist_ok=True)
        phase_entry = {
            "step": phase["step"],
            "folder": phase["folder"],
            "title": phase["title"],
            "files": []
        }

        # Création des fichiers par défaut s'ils n'existent pas
        for f in phase["files"]:
            file_path = os.path.join(phase_dir, f["filename"])
            if not os.path.exists(file_path):
                if f["type"] == "excel":
                    create_excel_template(file_path, f["title"], f["desc"], phase["title"])
                elif f["type"] == "word":
                    create_word_template(file_path, f["title"], f["desc"], phase["title"])

        # Scan dynamique de TOUS les fichiers réels présents dans le dossier
        for filename in sorted(os.listdir(phase_dir)):
            if filename.startswith(".") or filename in ("manifest.json", "manifest.js"):
                continue
            file_path = os.path.join(phase_dir, filename)
            if os.path.isdir(file_path):
                continue

            relative_path = f"documents_phases/{phase['folder']}/{filename}"
            ext = os.path.splitext(filename)[1].lower()
            file_type = "excel" if ext in (".xlsx", ".xls", ".csv") else "word" if ext in (".docx", ".doc") else "web" if ext in (".html", ".htm") else "pdf"

            # Recherche d'un titre et d'une description prédéfinis ou génération automatique propre
            matching_def = next((item for item in phase["files"] if item["filename"] == filename), None)
            if matching_def:
                label = matching_def["title"]
                desc = matching_def["desc"]
            else:
                # Nettoyage du nom pour affichage élégant
                clean_name = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ")
                label = clean_name
                desc = f"Document officiel de la phase {phase['step']} déposé dans {phase['folder']}."

            phase_entry["files"].append({
                "name": filename,
                "label": label,
                "desc": desc,
                "type": file_type,
                "path": relative_path
            })

        manifest_data.append(phase_entry)

    # Export JSON
    json_path = os.path.join(DOCS_DIR, "manifest.json")
    with open(json_path, "w", encoding="utf-8") as jf:
        json.dump(manifest_data, jf, indent=2, ensure_ascii=False)

    # Export JS pour exécution locale directe sans CORS
    js_path = os.path.join(DOCS_DIR, "manifest.js")
    with open(js_path, "w", encoding="utf-8") as jsf:
        jsf.write("/**\n * MANIFESTE DYNAMIQUE DES DOCUMENTS OFFICIELS PAR ÉTAPE\n * Mis à jour automatiquement ou modifiable librement.\n */\n")
        jsf.write("window.PHASE_DOCUMENTS_MANIFEST = " + json.dumps(manifest_data, indent=2, ensure_ascii=False) + ";\n")

    print(f"Successfully synced manifest with all phase folders in {DOCS_DIR}")

if __name__ == "__main__":
    main()
