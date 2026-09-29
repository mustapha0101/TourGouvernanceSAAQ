"""
Générateur du Document Officiel d'Expression des Besoins d'Affaires d'Investissement Québec (Fichier Démo)
Projet : SYNTH-IQ — Préqualification et Synthèse Intelligente des Dossiers de Financement

Ce document est conçu pour servir de démonstrateur d'ingestion automatique :
- Il présente un cas d'usage réaliste d'Investissement Québec
- Il comporte des points forts réels (valeur d'affaires, sponsor engagé)
- Il comporte des défis méthodologiques authentiques (hétérogénéité des données, PII Loi 25, formation)
- Il aboutit à un score moyen réaliste (~64-68 % / Homologation Conditionnelle)
- Il respecte la charte graphique : Titres Blanc sur Bleu Nuit (#002060), TrueType, tableaux propres.
"""

import os
import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

IQ_NAVY_HEX = "002060"      # Bleu Nuit Investissement Québec (#002060)
IQ_BLUE_HEX = "003366"      # Bleu Corporatif
IQ_BORDER_HEX = "CBD5E1"    # Bordure neutre
IQ_BG_LIGHT = "F8FAFC"      # Fond gris clair
IQ_AMBER_HEX = "D97706"     # Ambre

def set_cell_background(cell, fill_hex: str):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_section_header(doc, num: str, title: str):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    set_cell_background(c, IQ_NAVY_HEX)
    set_cell_margins(c, top=140, bottom=140, left=180, right=180)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{num} {title.upper()}")
    r.font.name = 'Calibri'
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # BLANC PUR SUR BLEU NUIT
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_demo_cahier_des_charges(output_path: str):
    doc = docx.Document()

    # Marges 2 cm (0.75 in)
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

        # En-tête & Pied de page
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("INVESTISSEMENT QUÉBEC • DIRECTION DU FINANCEMENT CORPORATIF • PROJET SYNTH-IQ")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Document d'expression des besoins d'affaires — Soumission officielle au Bureau de l'IA (Parcours 3)")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8.5)
        frun.font.italic = True
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Style de base
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # =========================================================================
    # BANDEAU DE TITRE OFFICIEL
    # =========================================================================
    t_title = doc.add_table(rows=1, cols=1)
    t_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_title = t_title.cell(0, 0)
    set_cell_background(c_title, IQ_NAVY_HEX)
    set_cell_margins(c_title, top=220, bottom=220, left=220, right=220)
    p_t = c_title.paragraphs[0]
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_t.add_run("INVESTISSEMENT QUÉBEC — GUICHET DE PRISE EN CHARGE DES INITIATIVES IA\n")
    r_inst.font.name = 'Calibri'
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0xBA, 0xE6, 0xFD)

    r_main = p_t.add_run("CAHIER D'EXPRESSION DES BESOINS D'AFFAIRES\nPROJET SYNTH-IQ (PRÉQUALIFICATION & ANALYSE DE CRÉDIT)")
    r_main.font.name = 'Calibri'
    r_main.font.size = Pt(15)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 1 : FICHE D'IDENTIFICATION DE L'INITIATIVE
    # =========================================================================
    add_section_header(doc, "1.", "Fiche d'Identification de l'Initiative IA")

    t_id = doc.add_table(rows=6, cols=2)
    t_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_id, color=IQ_BORDER_HEX)

    id_data = [
        ("Titre officiel du cas d'usage", "Préqualification et Synthèse Intelligente des Dossiers de Financement aux Entreprises (Projet SYNTH-IQ)"),
        ("Direction métier requérante", "Direction principale du Financement corporatif et Capital de risque"),
        ("Parrain d'affaires (Sponsor)", "Direction principale du Financement corporatif (Porteur Métier Délégué)"),
        ("Chargé de projet & Contact", "Équipe d'optimisation des processus de crédit — Direction du Financement"),
        ("Parcours officiel ciblé", "Parcours 3 : Cas d'usage IA & Accompagnement Stratégique Métier (Bureau de l'IA & Escouade IA)"),
        ("Type de technologie envisagée", "Système IA sur environnement fermé — Azure Cloud IQ (Architecture RAG hybride, OCR et LLM d'entreprise)")
    ]

    for r_i, (label, val) in enumerate(id_data):
        c0 = t_id.cell(r_i, 0)
        c1 = t_id.cell(r_i, 1)
        set_cell_background(c0, IQ_BG_LIGHT)
        set_cell_margins(c0, top=70, bottom=70, left=100, right=100)
        set_cell_margins(c1, top=70, bottom=70, left=100, right=100)
        r0 = c0.paragraphs[0].add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
        r1 = c1.paragraphs[0].add_run(val)
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # SECTION 2 : CONTEXTE D'AFFAIRES & OBJECTIF DE VALEUR MESURABLE
    # =========================================================================
    add_section_header(doc, "2.", "Contexte d'Affaires et Objectifs Opérationnels")

    p2_1 = doc.add_paragraph()
    p2_1.add_run("2.1 Description du Problème et Friction Métier Actuelle :\n").font.bold = True
    p2_1.add_run(
        "Chaque année, la Direction du Financement corporatif reçoit plus de 850 dossiers de demande de financement provenant de PME québécoises. "
        "Actuellement, un analyste financier consacre en moyenne entre 14 et 20 heures par dossier à extraire manuellement des données éparses : "
        "bilans financiers sur 3 exercices, états des résultats, organigrammes d'actionnariat, notes d'audit comptable et prévisions de trésorerie. "
        "Cette saisie manuelle engendre des goulots d'étranglement majeurs lors des périodes de fort volume (dépôts de bilans d'automne et de printemps), "
        "allonge les délais de réponse aux entrepreneurs et crée un risque d'erreur matérielle lors de la ressaisie dans nos modèles de ratio financier."
    )

    p2_2 = doc.add_paragraph()
    p2_2.add_run("2.2 Solution IA Proposée (Projet SYNTH-IQ) :\n").font.bold = True
    p2_2.add_run(
        "Déployer un assistant d'analyse documentaire sécurisé sur Azure Cloud IQ. Le système devra ingérer les documents d'affaires fournis par l'entreprise "
        "(fichiers PDF et tableurs), extraire les principaux postes comptables (chiffre d'affaires, BAIIA, dette nette, fonds de roulement), identifier les variations "
        "atypiques d'une année sur l'autre, et générer un premier brouillon de synthèse financière préqualifié destiné à l'analyste de crédit."
    )

    p2_3 = doc.add_paragraph()
    p2_3.add_run("2.3 Objectifs de Gains et Retour sur Investissement (ROI) :\n").font.bold = True
    p2_3.add_run(
        "• Gain de productivité : Réduction estimée de 4 à 6 heures de traitement par dossier de financement, soit plus de 4 000 heures réinvesties annuellement dans la relation d'affaires directe avec les entrepreneurs québécois.\n"
        "• Vélocité : Accélération de 30 % du délai d'émission du premier avis de préqualification financière.\n"
        "• Précision : Réduction de 85 % des erreurs de transcription comptable grâce à la traçabilité intégrale vers les documents sources originaux."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 3 : SOURCES DE DONNÉES & DÉFIS DE QUALITÉ IDENTIFIÉS
    # =========================================================================
    add_section_header(doc, "3.", "Corpus de Données et Défis de Qualité (Axe 2)")

    p3_1 = doc.add_paragraph()
    p3_1.add_run("3.1 Typologie des Documents Requis :\n").font.bold = True
    p3_1.add_run(
        "Le modèle analysera : (1) États financiers vérifiés ou avis au lecteur (PDF scannés ou natifs) ; (2) Déclarations fiscales provinciales et fédérales (CO-17, T2) ; "
        "(3) Plans d'affaires et prévisions de flux de trésorerie (Word, Excel) ; (4) Historique des transactions et encours dans les systèmes internes d'IQ (SharePoint, CRM de prêts)."
    )

    p3_2 = doc.add_paragraph()
    p3_2.add_run("⚠️ 3.2 Défis et Limites Réelles de Données (Facteur de Risque Moyen) :\n").font.bold = True
    p3_2.add_run(
        "• Hétérogénéité des formats : Plus de 35 % des documents reçus des PME sont des numérisations de qualité médiocre ou des documents manuscrits partiels, ce qui exigera une brique de reconnaissance optique de caractères (OCR) avancée.\n"
        "• Absence de standardisation comptable stricte : Les plans comptables varient d'une PME à l'autre; les métadonnées et libellés comptables ne sont pas homogènes.\n"
        "• Nécessité de nettoyage : Une phase de normalisation des données et de validation préalable des schémas comptables sera indispensable avant toute indexation vectorielle."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 4 : SENSIBILITÉ LOI 25 & RENSEIGNEMENTS PERSONNELS
    # =========================================================================
    add_section_header(doc, "4.", "Protection des Renseignements Personnels et Loi 25 (Axe 5)")

    p4_1 = doc.add_paragraph()
    p4_1.add_run("⚠️ 4.1 Présence de Données Nominatives et Sensibles (Point d'Attention Majeur) :\n").font.bold = True
    p4_1.add_run(
        "Bien que l'objet principal soit l'analyse d'entreprises (personnes morales), les dossiers de financement contiennent inévitablement des renseignements personnels sur des personnes physiques : "
        "noms, adresses résidentielles, participations au capital, déclarations de revenus des actionnaires garants et cotes de crédit personnelles des dirigeants."
    )

    p4_2 = doc.add_paragraph()
    p4_2.add_run("4.2 Exigences Légales et Engagements de Conformité d'IQ :\n").font.bold = True
    p4_2.add_run(
        "1. Déclenchement formel d'une Évaluation des Facteurs relatifs à la Vie Privée (ÉFVP) conformément à la Loi 25 du Québec, sous la supervision du Bureau de l'IA et du responsable de la protection des renseignements personnels (PRP).\n"
        "2. Environnement fermé avec Rétention Zéro Externe (Zero Data Retention) : Hébergement souverain sur tenant Azure Canada Central dédié d'IQ. Aucun apprentissage ni réentraînement du modèle externe sur les données clients.\n"
        "3. Anonymisation et masquage automatique des numéros d'assurance sociale (NAS) et données bancaires personnelles avant toute vectorisation."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 5 : SUPERVISION HUMAINE & GARDE-FOUS (HUMAN-IN-THE-LOOP)
    # =========================================================================
    add_section_header(doc, "5.", "Supervision Humaine Obligatoire et Niveau de Risque")

    p5_1 = doc.add_paragraph()
    p5_1.add_run("5.1 Règle Inviolable d'Assistance sans Décision Autonome (Human-in-the-Loop) :\n").font.bold = True
    p5_1.add_run(
        "Il est formellement interdit que le système SYNTH-IQ prenne une décision automatisée d'octroi, de rejet ou de tarification de financement. "
        "L'IA opère strictement en mode « Brouillon d'aide à la décision » : "
        "elle génère une synthèse annotée avec liens vers les sources factuelles. L'analyste de crédit titulaire conserve la responsabilité exclusive de la vérification, de la notation et de la recommandation transmise aux comités d'investissement."
    )

    p5_2 = doc.add_paragraph()
    p5_2.add_run("5.2 Gestion du Risque d'Hallucination et d'Exactitude Factuelle :\n").font.bold = True
    p5_2.add_run(
        "• Tout chiffre extrait ou calculé doit comporter une infobulle indiquant la page et la ligne exacte du document source d'origine.\n"
        "• Un score de confiance algorithmique (indicateur vert/jaune/rouge) avertira l'analyste dès qu'un document présente une incertitude de lecture supérieure à 10 %."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # SECTION 6 : ARCHITECTURE, RÉSILIENCE OPÉRATIONNELLE (RTO) & TI
    # =========================================================================
    add_section_header(doc, "6.", "Architecture Technique, Sécurité CSI et RTO (Axe 3 & 4)")

    t_tech = doc.add_table(rows=5, cols=2)
    t_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tech, color=IQ_BORDER_HEX)

    tech_data = [
        ("Architecture cible", "Cloud Azure IQ : Azure OpenAI Service (modèle privé dédié), Azure AI Search et Stockage Blob sécurisé."),
        ("Intégration écosystème", "Compatible SharePoint d'IQ, Teams, et connecteur vers le système d'octroi de prêts de la Société."),
        ("Délai de Reprise d'Activité (RTO)", "RTO ciblé à 48 heures. Le processus n'étant pas une infrastructure de paiement en temps réel, une interruption temporaire n'entraîne pas de paralysie critique."),
        ("Contrôle d'accès & CSI", "Gestion des accès basée sur les rôles (RBAC) héritée de l'annuaire corporatif Azure AD d'IQ avec journalisation d'audit complète."),
        ("Horizon de réalisation", "Preuve de concept (POC) en 8 semaines avec l'Escouade IA; déploiement pilote estimé à 6-8 mois.")
    ]

    for r_i, (k, v) in enumerate(tech_data):
        c0 = t_tech.cell(r_i, 0)
        c1 = t_tech.cell(r_i, 1)
        set_cell_background(c0, IQ_BG_LIGHT)
        set_cell_margins(c0, top=60, bottom=60, left=90, right=90)
        set_cell_margins(c1, top=60, bottom=60, left=90, right=90)
        r0 = c0.paragraphs[0].add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
        r1 = c1.paragraphs[0].add_run(v)
        r1.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # SECTION 7 : DEMANDE D'ACCOMPAGNEMENT AUPRÈS DU BUREAU DE L'IA
    # =========================================================================
    add_section_header(doc, "7.", "Demande d'Accompagnement et Attentes Métier")

    p7 = doc.add_paragraph()
    p7.add_run(
        "La Direction du Financement corporatif sollicite l'accompagnement officiel du Bureau de l'IA et de l'Escouade IA pour :\n"
        "1. Piloter l'atelier de cadrage méthodologique (Design Thinking IA) avec nos analystes et l'équipe Architecture TI.\n"
        "2. Coordonner l'Évaluation des Facteurs relatifs à la Vie Privée (ÉFVP) avec le délégué à la protection des données.\n"
        "3. Établir le cahier de tests et le jeu de données d'étalonnage (50 dossiers réels anonymisés) pour valider le taux d'exactitude comptable.\n"
        "4. Concevoir le plan de conduite du changement et former 35 analystes financiers à l'usage responsable de l'outil."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # =========================================================================
    # SECTION 8 : SIGNATURES ET ENGAGEMENT DU REQUÉRANT
    # =========================================================================
    add_section_header(doc, "8.", "Engagement du Parrain d'Affaires")

    t_s = doc.add_table(rows=2, cols=3)
    t_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_s, color=IQ_BORDER_HEX)

    headers_s = ["Instance Porteuse", "Représentant Désigné", "Date et Signature"]
    for c_i, h in enumerate(headers_s):
        c = t_s.cell(0, c_i)
        set_cell_background(c, IQ_NAVY_HEX)
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    r_cells = [
        "Direction principale du Financement corporatif",
        "Direction principale du Financement (Porteur Métier)",
        f"☑ Soumis le {datetime.date.today().strftime('%Y-%m-%d')}"
    ]
    for c_i, val in enumerate(r_cells):
        c = t_s.cell(1, c_i)
        set_cell_margins(c, top=60, bottom=60, left=80, right=80)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9)
        if c_i == 0:
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x20, 0x60)
        elif c_i == 2:
            r.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path

if __name__ == "__main__":
    out_dir = os.path.abspath(os.path.dirname(__file__))
    target_file = os.path.join(out_dir, "Cahier_des_Besoins_Demo_Precertification_Financement_IQ.docx")
    generate_demo_cahier_des_charges(target_file)
    print(f"Fichier Word généré avec succès : {target_file}")
