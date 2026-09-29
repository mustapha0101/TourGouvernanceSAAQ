"""
Générateur de livrables officiels Investissement Québec :
1. Grille Excel de Faisabilité complétée (.xlsx)
2. Mémo Décisionnel du Comité IA (.docx)
"""

import os
import shutil
import datetime
from typing import Dict, Any
import openpyxl
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

LOCAL_TEMPLATE_EXCEL = os.path.abspath(os.path.join(os.path.dirname(__file__), "templates", "Grille_faisabilite_cas_usage_IA.xlsx"))
LOCAL_TEMPLATE_DOCX = os.path.abspath(os.path.join(os.path.dirname(__file__), "templates", "00_Memo_Decisionnel_Transcription_v0.1.docx"))
FALLBACK_TEMPLATE_EXCEL = os.path.abspath(os.path.join(os.path.dirname(__file__), "../Analyse des initiatives /Grille_faisabilite_cas_usage_IA.xlsx"))
FALLBACK_TEMPLATE_DOCX = os.path.abspath(os.path.join(os.path.dirname(__file__), "../Analyse des initiatives /00_Memo_Decisionnel_Transcription_v0.1.docx"))

TEMPLATE_EXCEL = LOCAL_TEMPLATE_EXCEL if os.path.exists(LOCAL_TEMPLATE_EXCEL) else (FALLBACK_TEMPLATE_EXCEL if os.path.exists(FALLBACK_TEMPLATE_EXCEL) else LOCAL_TEMPLATE_EXCEL)
TEMPLATE_DOCX = LOCAL_TEMPLATE_DOCX if os.path.exists(LOCAL_TEMPLATE_DOCX) else (FALLBACK_TEMPLATE_DOCX if os.path.exists(FALLBACK_TEMPLATE_DOCX) else LOCAL_TEMPLATE_DOCX)
OUTPUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "generated_outputs"))

os.makedirs(OUTPUT_DIR, exist_ok=True)

# Lignes correspondantes aux 28 questions dans la feuille "Grille d'évaluation"
QUESTION_ROW_MAPPING = {
    # Axe 1 : Valeur
    1: 18, 2: 19, 3: 20, 4: 21, 5: 22,
    # Axe 2 : Données
    6: 26, 7: 27, 8: 28, 9: 29, 10: 30,
    # Axe 3 : Faisabilité
    11: 34, 12: 35, 13: 36, 14: 37, 15: 38,
    # Axe 4 : Effort
    16: 42, 17: 43, 18: 44, 19: 45,
    # Axe 5 : Risques
    20: 49, 21: 50, 22: 51, 23: 52, 24: 53,
    # Axe 6 : Adoption
    25: 57, 26: 58, 27: 59, 28: 60
}

# Mapping des 29 questions officielles de la feuille principale "Grille DD"
GRILLE_DD_ROW_MAPPING = {
    # Axe 1 : Valeur d'affaires et alignement stratégique (lignes 18 à 23)
    1: 18, 2: 19, 3: 20, 4: 21, 5: 22, 6: 23,
    # Axe 2 : Données disponibles et qualité (lignes 27 à 31)
    7: 27, 8: 28, 9: 29, 10: 30, 11: 31,
    # Axe 3 : Faisabilité technique (lignes 35 à 39)
    12: 35, 13: 36, 14: 37, 15: 38, 16: 39,
    # Axe 4 : Effort, coût et ressources (lignes 43 à 46)
    17: 43, 18: 44, 19: 45, 20: 46,
    # Axe 5 : Risques, conformité et Loi 25 (lignes 50 à 54)
    21: 50, 22: 51, 23: 52, 24: 53, 25: 54,
    # Axe 6 : Adoption et conduite du changement (lignes 58 à 61)
    26: 58, 27: 59, 28: 60, 29: 61
}

# Labels normalisés pour la colonne Réponse (C) de Grille DD (conformes aux balises IQ)
SCORE_LABELS_DD = {
    4: "4-Très élevé : Directement aligné / optimal",
    3: "3-Élevé : Conforme et maîtrisé",
    2: "2-Moyen : Partiel / mesures requises",
    1: "1-Faible : Faible maturité ou écart identifié",
    0: "0-Non applicable / Inexistant"
}

def generate_excel_feasibility(
    initiative_data: Dict[str, Any],
    evaluation_result: Dict[str, Any],
    file_prefix: str = "Grille_Faisabilite_IA"
) -> str:
    """
    Duplique et remplit intégralement le gabarit Excel officiel d'IQ avec les 28/29 critères calculés.
    Pré-complète à la fois la feuille active "Grille DD" et la feuille "Grille d'évaluation".
    """
    safe_title = "".join(c for c in initiative_data.get("title", "Projet") if c.isalnum() or c in (" ", "_", "-")).rstrip().replace(" ", "_")
    output_filename = f"{file_prefix}_{safe_title}_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    target_path = os.path.join(OUTPUT_DIR, output_filename)
    
    # Résolution résiliente du gabarit Excel officiel
    actual_template = None
    candidate_paths = [
        TEMPLATE_EXCEL,
        LOCAL_TEMPLATE_EXCEL,
        FALLBACK_TEMPLATE_EXCEL,
        os.path.join(os.getcwd(), "templates", "Grille_faisabilite_cas_usage_IA.xlsx")
    ]
    for cand in candidate_paths:
        if cand and os.path.exists(cand):
            actual_template = cand
            break

    if actual_template:
        shutil.copyfile(actual_template, target_path)
        wb = openpyxl.load_workbook(target_path)
    else:
        # Création de secours en cas d'absence imprévue du fichier physique
        wb = openpyxl.Workbook()
        ws_default = wb.active
        ws_default.title = "Grille DD"
        wb.create_sheet(title="Grille d'évaluation")
        wb.save(target_path)
    
    scores = evaluation_result.get("scores", {})
    justifications = evaluation_result.get("justifications", {})
    evaluator_name = evaluation_result.get("evaluator_name", "Bureau de l'IA & Escouade IA — Investissement Québec")
    today_str = datetime.date.today().strftime("%Y-%m-%d")
    
    # Mapping rigoureux des 29 critères de Grille DD vers les 28 critères officiels d'IQ
    DD_QUESTION_KEY_MAP = {
        1: "1",   # Problème clairement identifié
        2: "2",   # Besoin fréquent / fort volume
        3: "4",   # Alignement stratégique
        4: "3",   # Valeur attendue mesurable / qualité
        5: "5",   # Sponsor d'affaires / expérience
        6: "25",  # Utilisateurs / portée
        7: "6",   # Données nécessaires existent
        8: "7",   # Qualité des données
        9: "8",   # Volume de données
        10: "9",  # Données structurées / exploitables
        11: "10", # Droits d'utilisation
        12: "11", # Tâche adaptée à l'IA
        13: "12", # Solution existante sur le marché
        14: "13", # Intégration aux systèmes existants
        15: "14", # Niveau de précision / tolérance risque
        16: "15", # POC réalisable rapidement
        17: "16", # Effort proportionné
        18: "17", # Compétences disponibles
        19: "18", # Coûts maîtrisés
        20: "19", # Reproductibilité / extensibilité
        21: "20", # Respect Loi 25 PII
        22: "21", # Risques de biais / erreurs
        23: "22", # Humain dans la boucle (HITL)
        24: "23", # Traçabilité et auditabilité
        25: "24", # Sécurité de l'information (CSI)
        26: "25", # Adhésion utilisateurs finaux
        27: "26", # Intégration dans les flux de travail
        28: "27", # Plan de formation
        29: "28"  # Indicateurs de succès (KPI)
    }

    # -------------------------------------------------------------
    # 1. REMPLISSAGE INTÉGRAL DE LA FEUILLE PRINCIPALE "Grille DD"
    # -------------------------------------------------------------
    if "Grille DD" in wb.sheetnames:
        ws_dd = wb["Grille DD"]
        
        # Métadonnées en B4:D4 (fusionné), B5:D5, etc.
        ws_dd["B4"] = initiative_data.get("title", "Initiative IA IQ")
        ws_dd["B5"] = f"{initiative_data.get('direction', '')} / Porteur : {initiative_data.get('sponsor', 'Direction porteuse')}"
        ws_dd["B6"] = today_str
        ws_dd["B7"] = evaluator_name
        
        # Pré-remplissage des 29 critères dans Grille DD
        import iq_evaluation_engine
        for q_idx, row_num in GRILLE_DD_ROW_MAPPING.items():
            target_q_key = DD_QUESTION_KEY_MAP.get(q_idx, str(q_idx))
            sc = int(float(scores.get(target_q_key, 3)))
            sc = max(0, min(4, sc))
            label = SCORE_LABELS_DD.get(sc, f"{sc} - Conforme")
            
            # Colonne C (Réponse avec préfixe chiffré pour déclencher la formule LEFT(C,1)*1)
            ws_dd.cell(row=row_num, column=3, value=label)
            
            # Colonne E (Score chiffré direct pour compatibilité absolue tout tableur)
            ws_dd.cell(row=row_num, column=5, value=sc)
            
            # Colonne G (Justification de l'Analyste IA claire, contextuelle et motivée)
            raw_justif = str(justifications.get(target_q_key, "")).strip()
            if not raw_justif or "examiné et qualifié" in raw_justif.lower() or raw_justif.startswith("Critère #"):
                justif = iq_evaluation_engine.get_contextual_justification(target_q_key, initiative_data, sc)
            else:
                justif = raw_justif
            ws_dd.cell(row=row_num, column=7, value=justif)
            
        # Synthèse institutionnelle de l'Analyste IA en bas de la feuille Grille DD
        ws_dd.cell(row=72, column=1, value="SYNTHÈSE DE L'ANALYSTE EN GOUVERNANCE IA — BUREAU DE L'IA IQ")
        ws_dd.cell(row=72, column=1).font = openpyxl.styles.Font(bold=True, color="003366")
        ws_dd.cell(row=73, column=1, value=str(evaluation_result.get("executive_assessment", "Dossier conforme aux critères d'admissibilité.")))
        
        wb.active = ws_dd

    # -------------------------------------------------------------
    # 2. REMPLISSAGE DE LA FEUILLE "Grille d'évaluation"
    # -------------------------------------------------------------
    if "Grille d'évaluation" in wb.sheetnames:
        ws_eval = wb["Grille d'évaluation"]
        ws_eval["B4"] = initiative_data.get("title", "Initiative IA IQ")
        ws_eval["B5"] = f"{initiative_data.get('direction', '')} / Porteur : {initiative_data.get('sponsor', 'Non spécifié')}"
        ws_eval["B6"] = today_str
        ws_eval["B7"] = evaluator_name
        
        import iq_evaluation_engine
        for q_id, row_num in QUESTION_ROW_MAPPING.items():
            score_val = scores.get(str(q_id), scores.get(q_id, 3))
            try:
                ws_eval.cell(row=row_num, column=3, value=float(score_val))
            except Exception:
                ws_eval.cell(row=row_num, column=3, value=3)
                
            raw_justif = str(justifications.get(str(q_id), justifications.get(q_id, ""))).strip()
            if not raw_justif or "examiné et qualifié" in raw_justif.lower() or raw_justif.startswith("Critère #"):
                justif = iq_evaluation_engine.get_contextual_justification(str(q_id), initiative_data, score_val)
            else:
                justif = raw_justif
            ws_eval.cell(row=row_num, column=5, value=str(justif))
                
        ws_eval.cell(row=71, column=1, value="SYNTHÈSE DE L'ANALYSTE EN GOUVERNANCE IA — INVESTISSEMENT QUÉBEC")
        ws_eval.cell(row=71, column=1).font = openpyxl.styles.Font(bold=True, color="003366")
        ws_eval.cell(row=72, column=1, value=str(evaluation_result.get("executive_assessment", "Dossier conforme aux critères d'admissibilité.")))

    wb.save(target_path)
    return target_path

def clean_word_revisions_and_comments(doc: docx.Document):
    """
    Nettoie intégralement le document Word de toutes les marques de révision (Track Changes),
    commentaires résiduels, ratures et bulles de relecture du gabarit initial.
    Garantit un document final impeccable et officiel pour le Comité de gouvernance IA.
    """
    # 1. Supprimer toutes les suppressions suivies (w:del)
    for el in doc._element.xpath('//w:del'):
        p = el.getparent()
        if p is not None:
            p.remove(el)

    # 2. Accepter et déballer toutes les insertions suivies (w:ins)
    for el in doc._element.xpath('//w:ins'):
        p = el.getparent()
        if p is not None:
            for child in list(el):
                el.addprevious(child)
            p.remove(el)

    # 3. Supprimer toutes les balises de commentaires
    for tag in ['commentRangeStart', 'commentRangeEnd', 'commentReference']:
        for el in doc._element.xpath(f'//w:{tag}'):
            p = el.getparent()
            if p is not None:
                p.remove(el)

    # 4. Supprimer les marques d'historique de style et de propriétés
    for tag in ['rPrChange', 'pPrChange', 'tblPrChange', 'tcPrChange', 'sectPrChange']:
        for el in doc._element.xpath(f'//w:{tag}'):
            p = el.getparent()
            if p is not None:
                p.remove(el)

    # 5. Supprimer les relations vers les fichiers de commentaires du paquet Word
    try:
        for rel_id, rel in list(doc.part.rels.items()):
            rel_type = (rel.reltype or "").lower()
            if 'comments' in rel_type or 'people' in rel_type:
                del doc.part.rels[rel_id]
    except Exception as e:
        print("[Avertissement] Nettoyage part rels docx:", e)

def generate_word_memo(
    initiative_data: Dict[str, Any],
    evaluation_result: Dict[str, Any],
    memo_data: Dict[str, Any],
    file_prefix: str = "Memo_Decisionnel_Comite_IA"
) -> str:
    """
    Duplique et génère le Mémo Décisionnel exécutif d'IQ au format Word (.docx)
    selon la nouvelle charte multi-agents STORM (Titres Blanc sur Bleu, TOC hiérarchique, 4 perspectives).
    """
    import storm_memo_generator
    safe_title = "".join(c for c in initiative_data.get("title", "Projet") if c.isalnum() or c in (" ", "_", "-")).rstrip().replace(" ", "_")
    output_filename = f"{file_prefix}_{safe_title}_{datetime.date.today().strftime('%Y%m%d')}.docx"
    target_path = os.path.join(OUTPUT_DIR, output_filename)
    
    return storm_memo_generator.generate_storm_decision_memo(
        initiative_data=initiative_data,
        evaluation_result=evaluation_result,
        memo_data=memo_data,
        output_path=target_path
    )

